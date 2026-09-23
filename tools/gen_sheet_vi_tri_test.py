# -*- coding: utf-8 -*-
"""Sinh sheet "09. Vi tri file test" — yêu cầu chức năng ↔ đường dẫn tệp ↔ phương pháp.

    python tools/gen_sheet_vi_tri_test.py

Khác sheet 08 ở chỗ dùng để làm gì: sheet 08 trả lời "yêu cầu này đã được phủ tới mức
nào", còn sheet 09 trả lời "mở tệp nào để xem test của yêu cầu này, và test đó viết theo
kỹ thuật gì". Lúc bảo vệ, thầy hỏi "cho xem test của FR8.3" thì mở sheet này là có ngay
đường dẫn.

Đường dẫn tệp KHÔNG gõ tay: script tự dò trong src/test/java theo tên lớp, lớp nào không
tìm thấy thì dừng. Số test cũng đọc từ surefire. Phần viết tay duy nhất là ánh xạ
FR → lớp kiểm thử và nhãn phương pháp, dùng chung nguồn với sheet 08.
"""
import glob
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

XLSX = "BaoCao-KiemThu-CareerCompass.xlsx"
SHEET = "09. Vi tri file test"

C_HDR = "1F4E79"; C_TXT = "FFFFFFFF"; C_SEC = "DDEBF7"
C_HOPDEN = "FFF2CC"; C_HOPTRANG = "E2EFDA"; C_TICHHOP = "E4DFF5"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# Nhãn phương pháp — bám theo tên kỹ thuật trong bài giảng chương IV.
BVA = "Giá trị biên (BVA)"
EP = "Phân hoạch lớp tương đương"
TAG = "Gộp tag thành test case"
DT = "Bảng quyết định"
ST = "Chuyển đổi trạng thái"
UNIT = "Kiểm thử đơn vị (mock)"
SLICE = "Lát cắt web (@WebMvcTest)"
API = "Kiểm thử API (Postman)"
E2E = "Giao diện đầu-cuối (CodeceptJS)"

# (mã FR, tên, [(lớp kiểm thử, phương pháp)], ghi chú tầng khác)
FR = [
    ("FR1.1", "Hỏi đáp nghề nghiệp bằng ngôn ngữ tự nhiên",
     [("MentorServiceTest", UNIT), ("MentorControllerTest", SLICE)],
     "TC-MEN-01 → TC-MEN-05"),
    ("FR1.2", "Tích hợp mô hình ngôn ngữ lớn (thử lại + backoff)",
     [("LlmClientTest", UNIT)], "TC-MEN-04"),
    ("FR1.3", "Cá nhân hoá theo bảng điểm và GitHub",
     [("OnboardingFileSizeBvaTest", BVA),
      ("TranscriptFileDecisionTableTest", DT),
      ("OnboardingServiceTest", UNIT),
      ("TranscriptAnalysisServiceTest", UNIT),
      ("OnboardingControllerTest", SLICE)],
     "TC-ONB-05, TC-PR-04"),
    ("FR2.1", "Chọn vai trò nghề nghiệp mục tiêu",
     [("OnboardingStateTransitionTest", ST),
      ("OnboardingControllerTest", SLICE),
      ("UserProfileServiceTest", UNIT)],
     "TC-ONB-01 → TC-ONB-03"),
    ("FR2.2", "Sinh Skill Tree phân cấp theo mức ưu tiên",
     [("RoadmapServiceTest", UNIT), ("RoadmapControllerTest", SLICE),
      ("RoadmapPageControllerTest", SLICE)],
     "TC-RM-01, TC-RM-02"),
    ("FR2.3", "Mỗi node có tối thiểu 2 tài nguyên học",
     [("RoadmapServiceTest", UNIT)], "TC-RM-02"),
    ("FR2.4", "Đánh dấu hoàn thành và cập nhật tiến độ",
     [("ProgressDecisionTableTest", DT),
      ("ProgressStateTransitionTest", ST),
      ("ProgressServiceTest", UNIT),
      ("RoadmapControllerTest", SLICE)],
     "TC-RM-03 → TC-RM-07"),
    ("FR3.1", "Nhập/chọn kỹ năng hiện có",
     [("OnboardingControllerTest", SLICE), ("UserProfileServiceTest", UNIT)],
     "TC-ONB-06, TC-SG-02"),
    ("FR3.2", "So khớp kỹ năng với yêu cầu vai trò",
     [("SkillGapServiceTest", UNIT), ("SkillGapControllerTest", SLICE)],
     "TC-SG-01"),
    ("FR3.3", "Báo cáo trực quan, xuất PDF, ưu tiên học gấp",
     [("PdfServiceTest", UNIT), ("SkillGapControllerTest", SLICE),
      ("SkillGapPageControllerTest", SLICE)],
     "TC-SG-03 → TC-SG-09"),
    ("FR4.1", "Thu thập tin tuyển dụng định kỳ",
     [("ScraperServiceTest", BVA + " + " + UNIT)], "TC-MP-01"),
    ("FR4.2", "Phân tích tần suất từ khoá",
     [("KeywordAnalysisServiceTest", UNIT)], "TC-MP-01"),
    ("FR4.3", "Biểu đồ xu hướng tương tác",
     [("MarketPulseServiceTest", UNIT), ("MarketPulseControllerTest", SLICE)],
     "TC-MP-01"),
    ("FR5.1", "Liên kết GitHub và đồng bộ repository",
     [("PortfolioServiceTest", UNIT), ("PortfolioControllerTest", SLICE)],
     "TC-PF-01, TC-PF-02"),
    ("FR5.2", "Tóm tắt README bằng AI",
     [("PortfolioServiceTest", UNIT)], "TC-PF-02"),
    ("FR5.3", "Chia sẻ hồ sơ qua URL công khai",
     [("PortfolioControllerTest", SLICE)], "TC-PF-04, TC-SEC-06"),
    ("FR6.1", "Đăng ký, đăng nhập, khôi phục mật khẩu",
     [("RegisterStandardBvaTest", BVA),
      ("RegisterFormDTOBvaTest", BVA),
      ("PasswordPolicyBvaTest", BVA),
      ("TokenExpiryBvaTest", BVA),
      ("RegisterEquivalencePartitionTest", EP),
      ("RegisterTagCoverageTest", TAG),
      ("TokenValidityDecisionTableTest", DT),
      ("TokenStateTransitionTest", ST),
      ("AuthServiceTest", UNIT),
      ("PasswordResetServiceTest", UNIT),
      ("ProfileServiceTest", UNIT),
      ("CustomUserDetailsServiceTest", UNIT),
      ("AuthControllerTest", SLICE),
      ("HomeControllerTest", SLICE),
      ("ProfileControllerTest", SLICE)],
     "TC-AUTH-01 → TC-AUTH-17"),
    ("FR6.2", "Lưu lịch sử chat, đánh giá và tiến độ",
     [("DashboardServiceTest", UNIT), ("DashboardControllerTest", SLICE),
      ("ActivityLogServiceTest", UNIT)],
     "TC-DASH-01"),
    ("FR6.3", "Ép hoàn tất onboarding (quy tắc dẫn xuất)",
     [("OnboardingInterceptorTest", UNIT), ("OnboardingControllerTest", SLICE)],
     "TC-ONB-01, TC-SEC-01"),
    ("FR7.1", "Bảng điều khiển thống kê quản trị",
     [("AdminDashboardServiceTest", UNIT), ("AdminDashboardControllerTest", SLICE)],
     "TC-AD-01"),
    ("FR7.2", "Danh sách và tìm kiếm người dùng",
     [("AdminUserServiceTest", UNIT), ("AdminUserControllerTest", SLICE),
      ("UserAdminMapperTest", UNIT)],
     "TC-AD-02, TC-AD-03"),
    ("FR7.3", "Khoá/mở khoá, đổi vai trò, xoá người dùng",
     [("AdminUserServiceTest", UNIT), ("AdminUserControllerTest", SLICE)],
     "TC-AD-04 → TC-AD-07"),
    ("FR8.1", "Tạo, sửa, xoá template lộ trình",
     [("CounselorTemplateServiceTest", UNIT), ("CounselorTemplateControllerTest", SLICE)],
     "TC-CS-01 → TC-CS-03, TC-CS-11"),
    ("FR8.2", "Thêm/xoá node kỹ năng, đặt tầng và node cha",
     [("CounselorTemplateServiceTest", UNIT), ("CounselorTemplateControllerTest", SLICE)],
     "TC-CS-05, TC-CS-10"),
    ("FR8.3", "Gắn ≥ 2 tài liệu học cho mỗi node",
     [("CounselorTemplateServiceTest", UNIT), ("CounselorTemplateControllerTest", SLICE)],
     "TC-CS-08"),
    ("FR8.4", "Sắp xếp thứ tự ưu tiên node trong cùng tầng",
     [("CounselorTemplateServiceTest", UNIT), ("CounselorTemplateControllerTest", SLICE)],
     "TC-CS-07"),
]


def doc_surefire():
    kq = {}
    for p in glob.glob("target/surefire-reports/TEST-*.xml"):
        g = ET.parse(p).getroot().attrib
        kq[Path(p).stem[len("TEST-"):].rsplit(".", 1)[-1]] = int(g["tests"])
    if not kq:
        raise SystemExit("CHUA CHAY TEST: chay './mvnw clean test' truoc.")
    return kq


def dò_duong_dan():
    kq = {}
    for p in glob.glob("src/test/java/**/*Test.java", recursive=True):
        kq[Path(p).stem] = Path(p).as_posix()
    return kq


TEST = doc_surefire()
DUONG_DAN = dò_duong_dan()

for ma, _, ds, _ in FR:
    for lop, _pp in ds:
        if lop not in DUONG_DAN:
            raise SystemExit(f"{ma}: khong thay tep nguon cua lop '{lop}' trong src/test/java")
        if lop not in TEST:
            raise SystemExit(f"{ma}: lop '{lop}' khong co trong surefire-reports")

MAU_PP = {BVA: C_HOPDEN, EP: C_HOPDEN, TAG: C_HOPDEN, DT: C_HOPDEN, ST: C_HOPDEN,
          UNIT: C_HOPTRANG, SLICE: C_HOPTRANG}

wb = openpyxl.load_workbook(XLSX)
if SHEET in wb.sheetnames:
    wb.remove(wb[SHEET])
ws = wb.create_sheet(SHEET)
for cot, rong in zip("ABCDEF", (9, 34, 62, 27, 8, 24)):
    ws.column_dimensions[cot].width = rong


def o(r, c, v, *, bold=False, fill=None, center=False, size=10, color=None):
    x = ws.cell(row=r, column=c, value=v)
    x.font = Font(size=size, bold=bold, color=color)
    x.alignment = Alignment(vertical="center", wrap_text=True,
                            horizontal="center" if center else "left")
    x.border = BORDER
    if fill:
        x.fill = PatternFill("solid", fgColor=fill)


def thanh(r, t, *, fill=C_SEC, color=C_HDR, size=11, italic=False, cao=24):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    x = ws.cell(row=r, column=1, value=t)
    x.font = Font(bold=not italic, italic=italic, size=size, color=color)
    if fill:
        x.fill = PatternFill("solid", fgColor=fill)
    x.alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[r].height = cao
    return r + 1


r = 1
r = thanh(r, "VỊ TRÍ TỆP KIỂM THỬ THEO TỪNG YÊU CẦU CHỨC NĂNG", size=14, cao=30)
r = thanh(r, "Mỗi dòng là một lớp kiểm thử: mở đúng đường dẫn ở cột 3 là thấy mã nguồn. "
             "Cột phương pháp ghi kỹ thuật thiết kế test case đã dùng — ô vàng là kỹ "
             "thuật hộp đen (suy từ đặc tả), ô xanh là kỹ thuật hộp trắng (nhắm vào mã "
             "nguồn). Cột cuối dẫn sang mã ca kiểm thử tương ứng trong bộ Postman.",
           fill=None, color="595959", size=9, italic=True, cao=38)
r += 1

for i, h in enumerate(["Mã FR", "Yêu cầu", "Đường dẫn tệp kiểm thử",
                       "Phương pháp", "Số test", "Ca kiểm thử API"], 1):
    o(r, i, h, bold=True, fill=C_HDR, color=C_TXT, center=True)
ws.row_dimensions[r].height = 28
r += 1

tong = 0
for ma, ten, ds, tc in FR:
    dau = r
    for lop, pp in ds:
        o(r, 3, DUONG_DAN[lop], size=9)
        o(r, 4, pp, size=9, center=True, fill=MAU_PP.get(pp, C_TICHHOP))
        o(r, 5, TEST[lop], center=True)
        tong += TEST[lop]
        ws.row_dimensions[r].height = 18
        r += 1
    # Gộp ô mã FR / tên / ca kiểm thử cho cả nhóm dòng của yêu cầu đó.
    for cot, gia_tri, dam in ((1, ma, True), (2, ten, False), (6, tc, False)):
        o(dau, cot, gia_tri, bold=dam, center=(cot == 1),
          size=(10 if cot != 6 else 9))
        if r - dau > 1:
            ws.merge_cells(start_row=dau, start_column=cot, end_row=r - 1, end_column=cot)
        for rr in range(dau + 1, r):
            ws.cell(row=rr, column=cot).border = BORDER

o(r, 1, "TỔNG", bold=True, fill=C_SEC)
o(r, 2, f"{len(FR)} yêu cầu chức năng", bold=True, fill=C_SEC)
o(r, 3, f"{len(DUONG_DAN)} tệp kiểm thử trong src/test/java", bold=True, fill=C_SEC)
o(r, 4, "", fill=C_SEC)
o(r, 5, sum(TEST.values()), bold=True, fill=C_SEC, center=True)
o(r, 6, "", fill=C_SEC)
r += 2

r = thanh(r, "LƯU Ý KHI ĐỌC CỘT SỐ TEST: một lớp phục vụ nhiều yêu cầu sẽ xuất hiện ở "
             "nhiều dòng, nên cộng dồn cột này sẽ LỚN HƠN tổng thật. Dòng TỔNG lấy từ "
             f"surefire-reports: {sum(TEST.values())} test trên {len(TEST)} lớp.",
           fill=None, color="595959", size=9, italic=True, cao=28)
r = thanh(r, "CÁCH CHẠY MỘT LỚP BẤT KỲ:   ./mvnw test -Dtest=TênLớpTest     "
             "Chạy nhiều lớp:   ./mvnw test -Dtest='LớpA,LớpB'     "
             "Chạy một phương thức:   ./mvnw test -Dtest='LớpA#tênPhươngThức'",
           fill=None, color="595959", size=9, italic=True, cao=26)

wb.save(XLSX)
print(f"Da sinh sheet '{SHEET}'")
print(f"  {len(FR)} yeu cau chuc nang, {sum(len(d) for _, _, d, _ in FR)} dong tham chieu")
print(f"  {len(TEST)} lop kiem thu, {sum(TEST.values())} test")
