# -*- coding: utf-8 -*-
"""Thêm cột Status vào Bảng 5 của sheet 01.

    python tools/them_status_bang5.py

VẤN ĐỀ: Bảng 5 là bảng THIẾT KẾ TEST CASE duy nhất trong sheet 01 không có cột kết quả.
Bảng 2, Bảng 3 và Bảng 3b đều có "Kết quả thực tế" và "Status"; Bảng 5 thì chỉ có cột
"Test đã thi hành" ghi tên method, không nói method đó chạy ra xanh hay đỏ. Người đọc
nhìn 23 dòng mà không biết dòng nào đã đạt.

Cột 5, 6, 7 của Bảng 5 đang trống hẳn và vùng dữ liệu không có dải gộp nào, nên chỉ cần
ghi vào ô — không phải chèn cột, không phải dựng lại sheet.

Status lấy từ target/surefire-reports chứ không gõ tay: nếu một trong hai lớp thi hành
Bảng 5 có test đỏ thì script dừng, không cho báo cáo ghi PASS khống.

Chạy lại nhiều lần vẫn ra một bản — chỉ ghi đè đúng những ô đó.
"""
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

XLSX = "BaoCao-KiemThu-CareerCompass.xlsx"
SHEET = "01. BVA + Equiv Partition"
COT_STATUS = 7

C_HDR = "1F4E79"; C_TXT = "FFFFFFFF"; C_PASS = "C6EFCE"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# Hai lớp cùng thi hành Bảng 5: TC1–TC17 và TC18–TC23.
LOP = [("blackbox", "RegisterTagCoverageTest"), ("bva", "OnboardingFileSizeBvaTest")]

for pkg, cls in LOP:
    f = Path(f"target/surefire-reports/TEST-vn.uth.careercompass.{pkg}.{cls}.xml")
    if not f.exists():
        raise SystemExit(f"CHUA CHAY TEST: khong thay {f}. Chay 'mvnw test' truoc.")
    g = ET.parse(f).getroot().attrib
    if int(g["failures"]) or int(g["errors"]):
        raise SystemExit(f"{cls} dang do ({g['failures']} fail, {g['errors']} error) — "
                         "khong ghi PASS vao bao cao.")

wb = openpyxl.load_workbook(XLSX)
ws = wb[SHEET]

dau = next((r for r in range(1, ws.max_row + 1)
            if str(ws.cell(row=r, column=1).value or "").startswith("BẢNG 5")), None)
if dau is None:
    raise SystemExit("khong tim thay BANG 5 trong sheet 01")
hang_tieu_de = dau + 1
if str(ws.cell(row=hang_tieu_de, column=1).value or "") != "Test Case":
    raise SystemExit(f"dong {hang_tieu_de} khong phai hang tieu de cua Bang 5")

x = ws.cell(row=hang_tieu_de, column=COT_STATUS, value="Status")
x.font = Font(size=12, bold=True, color=C_TXT)
x.fill = PatternFill("solid", fgColor=C_HDR)
x.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center")
x.border = BORDER

n = 0
for r in range(hang_tieu_de + 1, ws.max_row + 1):
    a = ws.cell(row=r, column=1).value
    # Hết bảng khi gặp dòng không còn đánh số thứ tự test case.
    if not (isinstance(a, int) or str(a or "").isdigit()):
        break
    o = ws.cell(row=r, column=COT_STATUS, value="PASS")
    o.font = Font(size=10)
    o.fill = PatternFill("solid", fgColor=C_PASS)
    o.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center")
    o.border = BORDER
    n += 1

wb.save(XLSX)
print(f"Da them cot Status vao Bang 5 (dong tieu de {hang_tieu_de}), {n} dong ghi PASS")
