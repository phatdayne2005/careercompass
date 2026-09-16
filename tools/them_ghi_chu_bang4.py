# -*- coding: utf-8 -*-
"""Bổ sung lời giải thích vào ghi chú của Bảng 4, sheet 01.

    python tools/them_ghi_chu_bang4.py

VẤN ĐỀ: bảng tổng kết ghi RegisterEquivalencePartitionTest thi hành "Bảng 4 · V1–V4,
X1–X11". Đếm ra 4 + 11 = 15 nhãn, nhưng lớp đó chỉ có 14 test. Người đọc trừ ra thấy
thiếu một test mà không có chỗ nào giải thích.

Lý do thật: test của V3 dùng email "sinhvien@uth.edu.vn" — vừa đúng ĐỊNH DẠNG (V4) vừa
đúng ĐỘ DÀI (V3). Mà validateProperty kiểm mọi ràng buộc trên trường đó cùng lúc, nên
một test thoả cả hai lớp. 15 nhãn, 14 test.

Sửa TẠI CHỖ nội dung ô ghi chú sẵn có, không chèn dòng mới — sheet 01 có nhiều dải gộp,
mọi thao tác chèn dòng đều phải dựng lại cả sheet. Ở đây chỉ cần nối thêm câu nên sửa ô
là đủ và an toàn tuyệt đối.

Chạy lại nhiều lần vẫn ra một bản: nhận ra phần đã nối rồi thì không nối nữa.
"""
import openpyxl

XLSX = "BaoCao-KiemThu-CareerCompass.xlsx"
SHEET = "01. BVA + Equiv Partition"
NEO = "CODE KIỂM CHỨNG BẢNG NÀY"
DAU_HIEU = "VÌ SAO 15 NHÃN MÀ CHỈ 14 TEST"

THEM = (
    f"  ·  {DAU_HIEU}: bảng tổng kết ghi lớp này phủ V1–V4 và X1–X11, cộng lại là 15 "
    "nhãn, nhưng lớp chỉ có 14 test. Không thiếu test nào — test của V3 dùng email "
    "\"sinhvien@uth.edu.vn\", vừa đúng ĐỊNH DẠNG (V4) vừa đúng ĐỘ DÀI (V3). "
    "validateProperty kiểm mọi ràng buộc trên trường email cùng lúc nên một test thoả "
    "cả hai lớp hợp lệ. Bốn lớp không hợp lệ của email (X8, X9, X10, X11) thì phải tách "
    "riêng vì mỗi lớp vi phạm một ràng buộc khác nhau."
)

wb = openpyxl.load_workbook(XLSX)
ws = wb[SHEET]

hang = None
for r in range(1, ws.max_row + 1):
    if str(ws.cell(row=r, column=1).value or "").startswith(NEO):
        hang = r
        break
if hang is None:
    raise SystemExit(f"khong tim thay dong ghi chu bat dau bang '{NEO}'")

cu = str(ws.cell(row=hang, column=1).value or "")
if DAU_HIEU in cu:
    print(f"  dong {hang} da co loi giai thich, khong noi them")
else:
    ws.cell(row=hang, column=1, value=cu.rstrip() + THEM)
    ws.row_dimensions[hang] = ws.row_dimensions[hang]
    ws.row_dimensions[hang].height = 58
    print(f"  da noi loi giai thich vao ghi chu o dong {hang}")

wb.save(XLSX)
print(f"Da luu {XLSX}")
