# -*- coding: utf-8 -*-
"""Đóng dấu mốc thời gian đo lên hai sheet hộp trắng B3 và B4.

    python tools/dong_dau_moc_thoi_gian_whitebox.py

VẤN ĐỀ CẦN GIẢI QUYẾT: trong cùng một tệp báo cáo, sheet "00. Tong quan" ghi 399 test,
bao phủ dòng 78,6% và nhánh 72,5%; còn sheet B3 ghi 370 test, 75,5% và 67,5%. Người
chấm mở hai sheet cạnh nhau sẽ thấy số đá nhau và cho là làm ẩu.

Thực ra cả hai đều đúng, chỉ khác thời điểm đo: B1–B5 là bản chụp ngày 06/09/2026, sau
đó nhóm còn bổ sung thêm test nên các chỉ số tiếp tục tăng. Vấn đề chỉ nằm ở chỗ mốc
thời gian bị chôn trong một dòng phụ, không ai để ý.

CÁCH XỬ LÝ: chèn một dòng ghi chú nổi bật ngay đầu B3 và B4, nêu rõ ngày đo và dẫn sang
sheet 00 để xem số hiện tại. KHÔNG viết đè lên phân tích gốc — bảng B3 mô tả một phép
đo trước/sau có chủ đích của phần kiểm thử hộp trắng, viết số mới vào sẽ làm hỏng chính
mạch lập luận đó (mức tăng khi ấy là do test hộp trắng, không phải do test thêm về sau).

Chạy lại nhiều lần vẫn ra một bản: dòng ghi chú cũ bị gỡ trước khi chèn lại.
"""
import csv
import datetime as d

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

XLSX = "BaoCao-KiemThu-CareerCompass.xlsx"
CSV = "target/site/jacoco/jacoco.csv"
NGAY_DO_CU = "06/09/2026"
DAU_HIEU = "MỐC THỜI GIAN ĐO"          # để nhận ra dòng do script này chèn

C_CANH_BAO = "FFF2CC"
C_CHU = "7F6000"


def bao_phu():
    rows = list(csv.DictReader(open(CSV, encoding="utf-8")))

    def p(a, b):
        c = sum(int(r[a]) for r in rows)
        m = sum(int(r[b]) for r in rows)
        return 100 * c / (c + m)

    return p("LINE_COVERED", "LINE_MISSED"), p("BRANCH_COVERED", "BRANCH_MISSED")


def vn(x):
    return f"{x:.1f}".replace(".", ",")


LINE, BRANCH = bao_phu()
HOM_NAY = d.date.today().strftime("%d/%m/%Y")

GHI_CHU = (
    f"{DAU_HIEU}: các chỉ số ở sheet này được đo ngày {NGAY_DO_CU}, ngay sau khi bổ "
    f"sung phần kiểm thử hộp trắng — giữ nguyên để phép so sánh trước/sau còn ý nghĩa. "
    f"Nhóm tiếp tục bổ sung test sau mốc đó, nên số liệu HIỆN TẠI (đo {HOM_NAY}) cao "
    f"hơn: bao phủ dòng {vn(LINE)}%, bao phủ nhánh {vn(BRANCH)}%. "
    f"Xem sheet \"00. Tong quan\" để biết trạng thái mới nhất."
)

wb = openpyxl.load_workbook(XLSX)
for ten in ("B3_Truoc_va_Sau", "B4_Coverage_Per_Class"):
    ws = wb[ten]

    # Gỡ dòng ghi chú do lần chạy trước chèn, để chạy lại không xếp chồng.
    for r in range(1, min(ws.max_row, 8) + 1):
        if str(ws.cell(row=r, column=1).value or "").startswith(DAU_HIEU):
            ws.delete_rows(r)
            break

    # Chèn ngay dưới dòng tiêu đề và dòng mô tả nguồn của sheet.
    ws.insert_rows(3)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max(ws.max_column, 6))
    x = ws.cell(row=3, column=1, value=GHI_CHU)
    x.font = Font(bold=True, size=10, color=C_CHU)
    x.fill = PatternFill("solid", fgColor=C_CANH_BAO)
    x.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 44
    print(f"  da dong dau moc thoi gian len {ten}")

wb.save(XLSX)
print(f"Xong. Hien tai: dong {vn(LINE)}%, nhanh {vn(BRANCH)}%")
