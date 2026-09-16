# -*- coding: utf-8 -*-
"""Thêm một dòng ghi chú vào sheet 01 về lớp kiểm thử nằm ngoài bảng tổng kết.

    python tools/them_ghi_chu_sheet01.py

VÌ SAO CẦN: bảng tổng kết sheet 01 cộng ra 56 test từ bốn lớp. Nhưng trong mã nguồn còn
RegisterFormDTOBvaTest với 27 test cũng là kiểm thử giá trị biên. Việc loại nó ra là CÓ
CHỦ ĐÍCH và đã ghi lý do trong tools/gen_sheet01.py — nó chia mỗi trường thành một test
case riêng, trong khi slide 23 và 33 luôn trình bày test case là một bộ đầu vào đầy đủ.

Nhưng lý do đó chỉ nằm trong mã nguồn, người đọc báo cáo không thấy. Hệ quả: ai đếm số
test giá trị biên trong repo sẽ ra con số khác báo cáo mà không hiểu vì sao, và công sức
viết 27 test kia cũng không được nhắc tới ở đâu trong sheet.

KHÔNG dùng Worksheet.insert_rows(). Sheet 01 có nhiều dải gộp A..H; insert_rows dịch giá
trị ô nhưng KHÔNG dịch dải gộp, nên dải cũ trùm lên dòng mới và nuốt mất nội dung. Đã
gặp đúng lỗi này: dòng "max | 100 | 150 | 30" bị mất sạch trừ chữ "max". Thay vào đó đọc
cả sheet ra bộ nhớ rồi dựng lại, tính lại dải gộp theo vị trí mới — cùng cách mà
tools/cap_nhat_tong_quan.py đã dùng.

Chạy lại nhiều lần vẫn ra một bản: dòng cũ bị bỏ khi dựng lại.
"""
import copy
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

XLSX = "BaoCao-KiemThu-CareerCompass.xlsx"
SHEET = "01. BVA + Equiv Partition"
DAU_HIEU = "LỚP NẰM NGOÀI BẢNG TỔNG KẾT"

f = Path("target/surefire-reports/TEST-vn.uth.careercompass.bva.RegisterFormDTOBvaTest.xml")
if not f.exists():
    raise SystemExit(f"CHUA CHAY TEST: khong thay {f}. "
                     "Chay 'mvnw test -Dtest=RegisterFormDTOBvaTest' truoc.")
N = int(ET.parse(f).getroot().attrib["tests"])

GHI_CHU = (
    f"{DAU_HIEU}: mã nguồn còn RegisterFormDTOBvaTest ({N} test) cũng kiểm giá trị biên, "
    "nhưng KHÔNG tính vào bảng trên. Lý do: nó chia mỗi trường thành một test case riêng "
    "bằng validateProperty, trong khi slide 23 và 33 luôn trình bày test case là MỘT BỘ "
    "ĐẦU VÀO ĐẦY ĐỦ — Bảng 2, 3 và 5 của sheet này bám đúng mẫu đó. Các giá trị biên mà "
    "lớp kia kiểm đều đã được ba bảng nói trên phủ hết. Chạy riêng:  "
    f"mvnw test -Dtest=RegisterFormDTOBvaTest  →  {N} test."
)

wb = openpyxl.load_workbook(XLSX)
ws = wb[SHEET]

max_r, max_c = ws.max_row, ws.max_column
o_cu = [[ws.cell(row=r, column=c) for c in range(1, max_c + 1)]
        for r in range(1, max_r + 1)]
gop_cu = [(g.min_row, g.min_col, g.max_row, g.max_col) for g in ws.merged_cells.ranges]
cao_cu = {r: ws.row_dimensions[r].height for r in range(1, max_r + 1)
          if ws.row_dimensions[r].height}
rong = {k: v.width for k, v in ws.column_dimensions.items()}

# Bỏ dòng do lần chạy trước chèn, rồi tìm dòng neo để chèn lại ngay dưới.
giu = [r for r in range(1, max_r + 1)
       if not str(o_cu[r - 1][0].value or "").startswith(DAU_HIEU)]
neo = next((r for r in giu if str(o_cu[r - 1][0].value or "").startswith("Số ở cột")), None)
if neo is None:
    raise SystemExit("khong tim thay dong neo trong sheet 01")

anh_xa, moi, dong_moi = {}, 0, None
for i, r in enumerate(giu, start=1):
    anh_xa[r] = i + moi
    if r == neo:
        moi += 1
        dong_moi = i + moi

idx = wb.sheetnames.index(SHEET)
wb.remove(ws)
ws = wb.create_sheet(SHEET)
wb._sheets.remove(ws)
wb._sheets.insert(idx, ws)
for k, v in rong.items():
    ws.column_dimensions[k].width = v

for r_cu in giu:
    r_moi = anh_xa[r_cu]
    for c in range(1, max_c + 1):
        cu = o_cu[r_cu - 1][c - 1]
        x = ws.cell(row=r_moi, column=c, value=cu.value)
        x.font = copy.copy(cu.font)
        x.fill = copy.copy(cu.fill)
        x.alignment = copy.copy(cu.alignment)
        x.border = copy.copy(cu.border)
    if r_cu in cao_cu:
        ws.row_dimensions[r_moi].height = cao_cu[r_cu]

# Dải gộp tính lại theo vị trí mới; dải của dòng đã bỏ thì không còn đích để dịch.
for r1, c1, r2, c2 in gop_cu:
    if r1 in anh_xa and r2 in anh_xa:
        ws.merge_cells(start_row=anh_xa[r1], start_column=c1,
                       end_row=anh_xa[r2], end_column=c2)

ws.merge_cells(start_row=dong_moi, start_column=1, end_row=dong_moi, end_column=8)
x = ws.cell(row=dong_moi, column=1, value=GHI_CHU)
x.font = Font(italic=True, size=9, color="7F6000")
x.fill = PatternFill("solid", fgColor="FFF2CC")
x.alignment = Alignment(vertical="center", wrap_text=True)
ws.row_dimensions[dong_moi].height = 44

wb.save(XLSX)
print(f"Da them ghi chu vao sheet '{SHEET}', dong {dong_moi}")
print(f"  RegisterFormDTOBvaTest: {N} test (ngoai bang tong ket 56)")
