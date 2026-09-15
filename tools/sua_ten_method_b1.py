# -*- coding: utf-8 -*-
"""Sửa cột "Test Method" của bảng B1 sang tên phương thức CÓ THẬT trong mã nguồn.

    python tools/sua_ten_method_b1.py

VẤN ĐỀ: bảng B1 ghi 9 tên phương thức kiểm thử cho 9 đường cơ sở, nhưng KHÔNG tên nào
tồn tại trong repo (đã kiểm bằng grep toàn bộ src/, e2e/, tools/). Nếu giảng viên đọc
bảng rồi yêu cầu chạy thử một trong số đó, lệnh sẽ báo không tìm thấy test.

Các đường cơ sở đều ĐÃ được phủ thật, chỉ là tên phương thức khác. Bảng dưới đây ánh xạ
sang tên thật, mỗi ánh xạ đã đối chiếu bằng cách đọc thân phương thức chứ không đoán
theo tên.

Ba đường P5, P6, P7 của isNodeLocked() được phủ bởi một phương thức tham số hoá duy nhất
(lockExpression_coversConditionCombinations, 7 bộ tham số khớp đúng ma trận MC/DC ở bảng
B2), nên cột ghi kèm tên bộ tham số để biết dòng nào ứng với đường nào.

Sửa cả tệp nguồn BaoCao_Whitebox.xlsx, nếu không thì lần dựng lại báo cáo tiếp theo sẽ
ghi đè mất.
"""
import openpyxl

TEP = ["BaoCao-KiemThu-CareerCompass.xlsx", "BaoCao_Whitebox.xlsx"]
SHEET = "B1_CFG_va_Basis_Path"

# tên cũ (không tồn tại) -> tên thật trong RoadmapServiceTest.java
ANH_XA = {
    "calculatePercent_totalZero":
        "calculatePercent_whenTotalZero_returnsZero",
    "calculatePercent_normalCase":
        "calculatePercent_roundsToTwoDecimals",
    "isNodeLocked_tier1Null":
        "isNodeLocked_whenTierIsNull_treatsNodeAsTierOne",
    "isNodeLocked_explicitTier1":
        "isNodeLocked_whenTierOne_alwaysUnlocked",
    "isNodeLocked_tier2AllDone":
        "isNodeLocked_whenAllLowerTiersDone_returnsUnlocked",
    "isNodeLocked_tier2NotDone":
        "isNodeLocked_whenLowerTierNotAllDone_returnsLocked",
    "isNodeLocked_tier3Tier1NotDone":
        "lockExpression_coversConditionCombinations\n"
        "[tier 3 bị khóa khi tier 1 chưa xong]",
    "isNodeLocked_tier3Tier2NotDone":
        "lockExpression_coversConditionCombinations\n"
        "[tier 3 bị khóa khi tier 2 chưa xong]",
    "isNodeLocked_tier3AllDone":
        "lockExpression_coversConditionCombinations\n"
        "[tier 3 mở khi tier 1 và tier 2 đã xong]",
}

for tep in TEP:
    wb = openpyxl.load_workbook(tep)
    if SHEET not in wb.sheetnames:
        print(f"  bo qua {tep}: khong co sheet {SHEET}")
        continue
    ws = wb[SHEET]
    doi = 0
    for hang in ws.iter_rows():
        for o in hang:
            v = str(o.value or "").strip()
            if v in ANH_XA:
                o.value = ANH_XA[v]
                o.alignment = o.alignment.copy(wrap_text=True)
                doi += 1
    # Nới chiều cao các dòng có tên tham số hoá xuống hai dòng chữ.
    for r in range(11, min(ws.max_row, 22) + 1):
        if "lockExpression" in str(ws.cell(row=r, column=6).value or ""):
            ws.row_dimensions[r].height = 30
    wb.save(tep)
    print(f"  {tep}: doi {doi} ten phuong thuc")

print("\nDoi chieu lai bang lenh:")
print("  grep -c 'calculatePercent_whenTotalZero_returnsZero' "
      "src/test/java/vn/uth/careercompass/roadmap/service/RoadmapServiceTest.java")
